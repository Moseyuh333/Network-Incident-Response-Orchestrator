"""DOCX report builder for the Topic 09 submission."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
HEADER_FILL = "F2F4F7"
TABLE_WIDTHS = {
    "phase": [1.35, 2.25, 2.9],
    "mitre": [1.1, 1.9, 3.5],
    "actions": [0.6, 3.65, 1.4, 0.85],
}


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def _set_cell_margins(cell: Any, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _style_doc(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE

    sub = document.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    sub_run = sub.add_run(subtitle)
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    paragraph.add_run(text)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    hdr_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        hdr_cells[index].text = header
        hdr_cells[index].width = Inches(widths[index])
        _set_cell_shading(hdr_cells[index], HEADER_FILL)
        _set_cell_margins(hdr_cells[index])
        for paragraph in hdr_cells[index].paragraphs:
            for run in paragraph.runs:
                run.bold = True
        hdr_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            cells[index].width = Inches(widths[index])
            _set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def build_submission_docx(report: dict[str, Any], output_path: str | Path) -> Path:
    """Build the Vietnamese assignment report as a Word DOCX."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _style_doc(document)
    _add_title(
        document,
        "Topic 09 - Network Incident Response Orchestrator",
        "Bao cao thuc hanh: pipeline phan ung su co mang, song song hoa dieu tra va tao IR report.",
    )

    alert = report["alert"]
    classification = report["stage_2_analysis"]["incident_classification"]
    document.add_heading("I. Tong quan", level=1)
    document.add_paragraph(
        "De tai xay dung mot pipeline phan ung su co mang duoc kich hoat boi mot "
        "alert. He thong gom cac tac vu thu thap doc lap, phan loai su co, anh xa "
        "MITRE ATT&CK va de xuat cac buoc containment an toan."
    )
    _add_bullet(document, "Muc tieu: rut ngan thoi gian triage bang song song hoa recon, log collection va PCAP feature extraction.")
    _add_bullet(document, "Pham vi: defensive lab, khong thuc hien hanh dong chan/mat mang that.")
    _add_bullet(document, "Phan cong mau: SOC L1 nhan alert, SOC L2 dieu phoi pipeline, DFIR thu thap bang chung, Network Security containment.")

    document.add_heading("II. Noi dung", level=1)
    document.add_heading("1. Pipeline thuc thi", level=2)
    _add_table(
        document,
        ["Phase", "Kieu chay", "Mo ta"],
        [
            ["Stage 0", "Tuan tu", "Normalize alert, gan alert_id, source, destination, severity."],
            ["Stage 1", "Song song", "Recon host, thu thap log, trich xuat PCAP features."],
            ["Stage 2", "Song song", "Classifier co trong so va embedding-style profile scoring."],
            ["Stage 3", "Tuan tu", "Map MITRE ATT&CK va sinh containment plan."],
            ["Stage 4", "Tuan tu", "Xuat JSON, Markdown, DOCX va zip project."],
        ],
        TABLE_WIDTHS["phase"],
    )
    document.add_paragraph(
        "Loi ich song song hoa: recon, log va PCAP khong phu thuoc lan nhau nen co "
        "the chay cung luc. Sau khi Stage 1 hoan tat, classifier va profile scoring "
        "cung co the chay doc lap tren cung bo evidence."
    )

    document.add_heading("2. Scripts, skills, agent va chain", level=2)
    _add_bullet(document, "Python package: app/orchestrator gom collectors, classifier, MITRE mapper, reporting va pipeline.")
    _add_bullet(document, "Script chay demo: scripts/run_demo.py tao .pi artifacts, chay pipeline va sinh bao cao.")
    _add_bullet(document, "Skill: .pi/skills/network-ir-orchestrator/SKILL.md mo ta workflow IR.")
    _add_bullet(document, "Agent: .pi/agents/incident-orchestrator-agent.md dieu phoi cac stage.")
    _add_bullet(document, "Chain: .pi/chains/network-ir-chain.md mo ta luong alert -> collect -> classify -> report.")
    _add_bullet(document, "Packages can thiet: python-docx cho DOCX; pytest/unittest de kiem thu; FastAPI stack trong pyproject cho mo rong API.")

    document.add_heading("3. Ket qua phan tich alert mau", level=2)
    _add_table(
        document,
        ["Truong", "Gia tri"],
        [
            ["Alert ID", str(alert["alert_id"])],
            ["Nguon", str(alert["source_ip"])],
            ["Dich", str(alert["destination_ip"])],
            ["Loai su co", str(classification["incident_type"])],
            ["Severity", str(classification["severity"])],
            ["Confidence", str(classification["confidence"])],
        ],
        [1.7, 4.8],
    )

    document.add_heading("4. MITRE ATT&CK mapping", level=2)
    _add_table(
        document,
        ["ID", "Tactic", "Technique"],
        [
            [item["technique_id"], item["tactic"], item["technique"]]
            for item in report["mitre_attack"]
        ],
        TABLE_WIDTHS["mitre"],
    )

    document.add_heading("5. Containment steps", level=2)
    _add_table(
        document,
        ["#", "Action", "Owner", "Mode"],
        [
            [
                str(item["step"]),
                item["action"],
                item["owner"],
                item["mode"],
            ]
            for item in report["containment_plan"]
        ],
        TABLE_WIDTHS["actions"],
    )

    document.add_heading("6. Huong dan su dung", level=2)
    _add_bullet(document, "Chay demo: python scripts/run_demo.py")
    _add_bullet(document, "Chay test: python -m unittest discover -s tests")
    _add_bullet(document, "Prompt mau: 'Phan tich alert ALERT-09-001 va tao IR report voi MITRE mapping va containment plan.'")

    document.add_heading("III. Danh gia ket qua va ket luan", level=1)
    document.add_paragraph(
        "Pipeline da tao duoc ket qua co cau truc gom JSON, Markdown va DOCX. Ket qua "
        "the hien duoc loi ich cua song song hoa trong incident response: cac tac vu "
        "thu thap bang chung doc lap duoc gop lai nhanh, sau do moi chuyen sang phan "
        "tich va ra quyet dinh containment."
    )
    document.add_paragraph(
        "Gioi han: classifier hien tai la offline deterministic model phuc vu lab. Khi "
        "trien khai thuc te co the thay bang ML model hoac LLM co guardrails va approval gate."
    )

    document.add_heading("IV. Tai lieu tham khao", level=1)
    _add_bullet(document, "MITRE ATT&CK: https://attack.mitre.org/")
    _add_bullet(document, "NIST SP 800-61 Rev. 2 Computer Security Incident Handling Guide.")
    _add_bullet(document, "FastAPI, python-docx va Python concurrent.futures documentation.")

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Topic 09 - Network IR Orchestrator")

    document.save(output)
    return output
